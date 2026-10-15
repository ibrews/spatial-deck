#!/usr/bin/env python3
"""Editable PowerPoint of the PMRE talk, built from talk/sections.js (not screenshots).

Every title, quote, tagline and bullet is a real text box in a font every show machine has (Arial), so the
organizer's production team can reformat for their non-standard LED wall. Images are placed by aspect ratio;
motion slides use the rendered poster frame and name their MP4 (media/motion/mp4/) in the notes — PMRE asked for
video as separate files, not embedded. Full presenter notes go in the notes pane.
    python3 tools/export_talk_pptx.py --out ~/Archives/pmre2026-2026-10/exports/pmre2026-xr-chitect-editable.pptx
"""
import argparse, html, json, re, subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
FONT = "Arial"
BG, TEXT, DIM = RGBColor(0x06, 0x08, 0x10), RGBColor(0xF0, 0xF4, 0xFF), RGBColor(0x9A, 0xA0, 0xB4)
ACC = {"teal": RGBColor(0x00, 0xD4, 0xFF), "purple": RGBColor(0xA7, 0x8B, 0xFA), "amber": RGBColor(0xF5, 0x9E, 0x0B), "rose": RGBColor(0xF4, 0x3F, 0x5E)}
W, H = Inches(13.333), Inches(7.5)

def load_talk():
    js = (ROOT / "talk/sections.js").read_text(encoding="utf-8")
    out = subprocess.run(["node", "-e", js + ";process.stdout.write(JSON.stringify({TALK,SECTIONS,BONUS}))"], capture_output=True, text=True, check=True).stdout
    return json.loads(out)

def clean(t):
    t = re.sub(r"<br\s*/?>\s*(<br\s*/?>)?", "\n", t or "")
    return html.unescape(re.sub(r"<[^>]+>", "", t)).strip()

def text(slide, x, y, w, h, s, size, color=TEXT, bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, line in enumerate(clean(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        if spacing: p.space_after = Pt(spacing)
        r = p.add_run(); r.text = line; f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color
    return tb

def bullets(slide, x, y, w, h, items, accent):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, b in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(14)
        r = p.add_run(); r.text = "—  "; r.font.name = FONT; r.font.size = Pt(18); r.font.color.rgb = accent; r.font.bold = True
        r = p.add_run(); r.text = clean(b); r.font.name = FONT; r.font.size = Pt(18); r.font.color.rgb = TEXT

def picture(slide, path, x, y, w, h, cover=False):
    p = ROOT / path
    if not p.exists() or p.suffix.lower() in (".mp4", ".webm", ".svg", ".html"):
        return None
    iw, ih = Image.open(p).size; r = iw / ih
    if cover:
        pic = slide.shapes.add_picture(str(p), x, y, w, h)
        box = w / h
        if r > box: c = (1 - box / r) / 2; pic.crop_left = pic.crop_right = c
        else: c = (1 - r / box) / 2; pic.crop_top = pic.crop_bottom = c
        return pic
    if r > w / h: pw, ph = w, int(w / r)
    else: pw, ph = int(h * r), h
    return slide.shapes.add_picture(str(p), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)

def new(prs, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    s.notes_slide.notes_text_frame.text = notes or ""
    return s

def big_size(t):
    n = max(len(l) for l in clean(t).split("\n"))
    return 66 if n <= 20 else 54 if n <= 27 else 44 if n <= 34 else 36

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True)
    ap.add_argument("--pick", action="append", default=[], help="group=choice, e.g. --pick open=question --pick cut=30 (defaults from TALK.options)")
    a = ap.parse_args()
    d = load_talk(); T, S, B = d["TALK"], d["SECTIONS"], d["BONUS"]
    pick = {g: o["default"] for g, o in T.get("options", {}).items()}
    for kv in a.pick:
        g, c = kv.split("=", 1); assert g in pick and c in T["options"][g]["choices"], f"unknown option {kv}"; pick[g] = c
    cut = None if pick.get("cut", "full") == "full" else int(pick["cut"])
    def shown(c):
        if c.get("opt"):
            g, ch = c["opt"].split(":")
            if pick.get(g) not in ch.split("|"): return False
        return cut is None or (c.get("keep") or 99) <= cut
    th = T.get("thesisAlts", {}).get(pick.get("thesis")); cl = T.get("closeAlts", {}).get(pick.get("close"))
    if th: B = {**B, "title": th["title"], "tagline": th["tagline"]}
    if cl: T = {**T, "closeTitle": cl["title"], "closeSub": cl["sub"]}
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    videos = []
    s = new(prs, T["coverNotes"])
    text(s, Inches(1), Inches(2.0), Inches(11.3), Inches(.5), T["coverEvent"], 14, ACC["teal"], bold=True, align=PP_ALIGN.CENTER)
    text(s, Inches(1), Inches(2.6), Inches(11.3), Inches(1.4), T["coverTitle"], 60, bold=True, align=PP_ALIGN.CENTER)
    text(s, Inches(1), Inches(4.0), Inches(11.3), Inches(.7), T["coverSub"], 26, DIM, align=PP_ALIGN.CENTER)
    text(s, Inches(1), Inches(5.2), Inches(11.3), Inches(.9), T["speaker"] + "\n" + T["speakerCo"], 18, TEXT, align=PP_ALIGN.CENTER)
    for si, sec in enumerate(S):
        ac = ACC.get(sec["accent"], ACC["teal"]); L = sec["lesson"]
        s = new(prs, L.get("notes"))
        text(s, Inches(1), Inches(1.4), Inches(11.3), Inches(.5), f"{sec['year']}  ·  {si + 1} / {len(S)}", 16, ac, bold=True, align=PP_ALIGN.CENTER)
        text(s, Inches(1), Inches(2.0), Inches(11.3), Inches(2.2), L["title"], 54, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(2), Inches(4.4), Inches(9.3), Inches(1.9), L["tagline"], 20, DIM, align=PP_ALIGN.CENTER, spacing=8)
        if L.get("tags"): text(s, Inches(1), Inches(6.5), Inches(11.3), Inches(.4), L["tags"], 12, ac, align=PP_ALIGN.CENTER)
        for c in sec["cases"]:
            if not shown(c): continue
            s = new(prs, c.get("notes")); lay = c.get("layout")
            if lay == "full" and c.get("motion"):
                m = c["motion"]; poster = f"media/motion/mp4/{m}-poster.png"
                picture(s, poster, 0, 0, W, H)
                s.notes_slide.notes_text_frame.text = f"▶ VIDEO: play media/motion/mp4/{m}.mp4 full-screen on this slide (submitted as a separate file).\n\n" + (c.get("notes") or "")
                videos.append((len(prs.slides), m))
            elif lay == "big":
                if c.get("title"): text(s, Inches(1), Inches(1.0), Inches(11.3), Inches(.6), c["title"], 22, DIM, align=PP_ALIGN.CENTER)
                text(s, Inches(.8), Inches(1.8), Inches(11.7), Inches(3.8), c["bigText"], big_size(c["bigText"]), TEXT, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                if c.get("bigCaption"): text(s, Inches(1.5), Inches(5.9), Inches(10.3), Inches(.9), c["bigCaption"], 16, DIM, italic=True, align=PP_ALIGN.CENTER)
            elif lay == "grid":
                ims = c["images"]; n = len(ims); cols = 2 if n == 4 else n; rows = 2 if n == 4 else 1
                gw, gh = W // cols, (H - Inches(1.1)) // rows
                for i, it in enumerate(ims):
                    x, y = (i % cols) * gw, (i // cols) * gh
                    picture(s, it["src"] if isinstance(it, dict) else it, x + Inches(.05), y + Inches(.05), gw - Inches(.1), gh - Inches(.1), cover=bool(c.get("gridCover")))
                    if isinstance(it, dict) and it.get("label"): text(s, x + Inches(.15), y + Inches(.1), gw - Inches(.3), Inches(.4), it["label"], 12, TEXT, bold=True)
                text(s, Inches(.5), H - Inches(1.05), Inches(8), Inches(.5), c.get("title", ""), 26, TEXT, bold=True)
                text(s, Inches(.5), H - Inches(.55), Inches(10), Inches(.4), c.get("subtitle", ""), 14, DIM)
            else:
                picture(s, c.get("img", ""), Inches(.6), Inches(1.2), Inches(6.2), Inches(5.1))
                text(s, Inches(7.2), Inches(1.5), Inches(5.6), Inches(.4), sec["year"] + "  ·  " + L.get("short", ""), 12, ac, bold=True)
                text(s, Inches(7.2), Inches(1.9), Inches(5.6), Inches(1.3), c.get("title", ""), 32, bold=True)
                text(s, Inches(7.2), Inches(3.1), Inches(5.6), Inches(.6), c.get("subtitle", ""), 16, DIM, italic=True)
                bullets(s, Inches(7.2), Inches(3.8), Inches(5.6), Inches(3.0), c.get("bullets", []), ac)
    s = new(prs, B.get("notes"))
    text(s, Inches(1), Inches(1.4), Inches(11.3), Inches(.5), B["year"], 16, ACC["amber"], bold=True, align=PP_ALIGN.CENTER)
    text(s, Inches(.8), Inches(2.0), Inches(11.7), Inches(2.2), B["title"], 50, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(2), Inches(4.4), Inches(9.3), Inches(1.9), B["tagline"], 20, DIM, align=PP_ALIGN.CENTER, spacing=8)
    s = new(prs, T["closeNotes"])
    text(s, Inches(1), Inches(1.5), Inches(11.3), Inches(.5), T["closeEyebrow"].upper(), 14, ACC["amber"], bold=True, align=PP_ALIGN.CENTER)
    text(s, Inches(1), Inches(2.0), Inches(11.3), Inches(1.3), T["closeTitle"], 60, bold=True, align=PP_ALIGN.CENTER)
    text(s, Inches(1), Inches(3.3), Inches(11.3), Inches(.7), T["closeSub"], 22, DIM, align=PP_ALIGN.CENTER)
    picture(s, "talk/qr.png", Inches(5.67), Inches(4.3), Inches(2), Inches(2))
    text(s, Inches(1), Inches(6.4), Inches(11.3), Inches(.5), T["web"], 16, ACC["teal"], align=PP_ALIGN.CENTER)
    out = Path(a.out).expanduser(); out.parent.mkdir(parents=True, exist_ok=True); prs.save(out)
    print(f"{out}: {len(prs.slides)} slides, picks={pick}, fonts={FONT}; video slides: " + ", ".join(f"#{n} {m}.mp4" for n, m in videos))

if __name__ == "__main__":
    main()
